"""Reading a Word format, and binding into it.

The format is the product's surface for a firm: they open a `.docx` in Word, change how a
finding looks, and every report changes. So the contract tested here is the one an author
relies on — which bracketed tokens mean what, that the worked example is the per-finding
shape, and that a mistake in a format is an error naming the fix rather than a gap in a
document a buyer reads.
"""

from __future__ import annotations

from typing import Any

import pytest

from inspection_report import sample
from inspection_report.domain import catalogue
from inspection_report.render import report as render_report
from inspection_report.templates import binding, docx_html

SYSTEMS = [s.name for s in catalogue.systems()]


def a_format(
    *,
    with_photos: bool = True,
    example_extra: str = "",
    system_extra: str = "",
    head: str = "<p>[firm name] — [property address] — [inspector] — [date of inspection]</p>",
) -> str:
    photo = "<p>[photograph]</p><p><em>[caption]</em></p>" if with_photos else ""
    return (
        head
        + "<h1>How each finding is recorded</h1>"
        + "<p><strong>[severity label]: [location]</strong></p>"
        + "<p>[observation]</p>"
        + example_extra
        + "<p><strong>Recommended next step: </strong>[recommendation]</p>"
        + photo
        + "<h1>What was inspected</h1><p>[systems inspected]</p>"
        + "".join(
            f"<h1>{name}</h1><p><em>blurb</em></p>{system_extra}<p>[findings]</p>"
            for name in SYSTEMS
        )
    )


class TestReadingAFormat:
    def test_the_worked_example_becomes_the_finding_shape(self) -> None:
        fmt = binding.read_format(a_format(), SYSTEMS)
        assert fmt.finding_shape.carries_photos
        assert len(fmt.finding_shape.photo) == 2, "the photograph and its caption repeat together"
        assert len(fmt.system_slots) == len(SYSTEMS)

    def test_the_observation_paragraph_is_marked_for_the_rewrite_pass(self) -> None:
        """The AI is pointed at a class, and `class` is what survives an upload."""
        fmt = binding.read_format(a_format(), SYSTEMS)
        marked = [raw for raw in fmt.finding_shape.before if binding.NOTE_CLASS in raw]
        assert len(marked) == 1
        assert "[observation]" in marked[0]

    def test_a_format_with_no_photograph_in_its_example_carries_none(self) -> None:
        fmt = binding.read_format(a_format(with_photos=False), SYSTEMS)
        assert not fmt.finding_shape.carries_photos

    def test_a_mistyped_token_is_refused_and_the_vocabulary_is_named(self) -> None:
        with pytest.raises(binding.TemplateError) as exc:
            binding.read_format(a_format(head="<p>[firm nmae]</p>"), SYSTEMS)
        assert "firm nmae" in str(exc.value)
        assert "severity label" in str(exc.value), "the error should teach the vocabulary"

    def test_an_interrupted_worked_example_is_refused(self) -> None:
        """The example is repeated per finding, so unrelated prose must not be inside it."""
        with pytest.raises(binding.TemplateError) as exc:
            binding.read_format(a_format(example_extra="<p>A note to the author.</p>"), SYSTEMS)
        assert "interrupted" in str(exc.value)

    def test_a_missing_system_is_named(self) -> None:
        html = a_format().replace(f"<h1>{SYSTEMS[2]}</h1>", "<h1>Something else</h1>")
        with pytest.raises(binding.TemplateError) as exc:
            binding.read_format(html, SYSTEMS)
        assert SYSTEMS[2] in str(exc.value)


class TestBinding:
    def test_the_example_section_never_reaches_the_report(self) -> None:
        html = render_report.render(sample.sample_inspection(), a_format())
        assert "How each finding is recorded" not in html
        assert "[observation]" not in html

    def test_the_photograph_run_repeats_once_per_photograph(self) -> None:
        inspection = sample.sample_inspection()
        for finding in inspection.findings:
            for photo in finding.photos:
                photo.remote_url = f"https://example.invalid/{photo.sha256[:8]}.png"
        html = render_report.render(inspection, a_format())
        assert html.count("<img ") == inspection.photo_count()

    def test_a_photograph_that_never_uploaded_is_left_out(self) -> None:
        """A broken image in a document someone prints is worse than no image."""
        html = render_report.render(sample.sample_inspection(), a_format())
        assert "<img " not in html

    def test_a_system_with_nothing_under_it_says_so_in_the_formats_own_paragraph(self) -> None:
        inspection = sample.sample_inspection()
        inspection.findings = [f for f in inspection.findings if f.system_key == "roof"]
        html = render_report.render(inspection, a_format())
        assert html.count(render_report.NOTHING_OBSERVED) == len(SYSTEMS) - 1

    def test_a_token_nobody_fills_is_an_error_rather_than_a_gap(self) -> None:
        html = a_format().replace("<p>[systems inspected]</p>", "<p>[system summary]</p>")
        # [system summary] outside a system section has no value to take, so it survives to
        # the end — which must fail loudly rather than print an empty line to a buyer.
        with pytest.raises(binding.TemplateError) as exc:
            render_report.render(sample.sample_inspection(), html)
        assert "system summary" in str(exc.value)

    def test_an_inspectors_angle_brackets_stay_text(self) -> None:
        inspection = sample.sample_inspection()
        inspection.findings[0].observation = "gap <2in at flashing"
        html = render_report.render(inspection, a_format())
        assert "&lt;2in" in html


class TestStampClass:
    def test_it_adds_a_class_where_there_is_none(self) -> None:
        assert binding.stamp_class("<p>x</p>", "note") == '<p class="note">x</p>'

    def test_it_keeps_the_classes_already_there(self) -> None:
        out = binding.stamp_class('<p class="lead">x</p>', "note")
        assert 'class="lead note"' in out

    def test_it_does_not_add_the_same_class_twice(self) -> None:
        once = binding.stamp_class("<p>x</p>", "note")
        assert binding.stamp_class(once, "note") == once

    def test_it_leaves_other_attributes_alone(self) -> None:
        out = binding.stamp_class('<p style="text-align:center">x</p>', "note")
        assert "text-align:center" in out
        assert 'class="note"' in out


class TestTheConverterMatchesWhatTheServiceReturns:
    """The offline converter exists so the fake and the fallback see what live sees.

    A converter that produced something richer would let a format work with no key and fail
    against the service, which is the exact failure this build has been bitten by before.
    """

    def _doc(self) -> Any:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        doc = Document()
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("ACME")
        run.bold = True
        run.font.size = Pt(18)
        doc.add_heading("A section", level=1)
        para = doc.add_paragraph()
        para.add_run("Label: ").bold = True
        para.add_run("value")
        doc.add_paragraph()  # spacing, the way Word documents carry it
        doc.add_paragraph("first", style="List Bullet")
        doc.add_paragraph("second", style="List Bullet")
        doc.add_paragraph().add_run("in italics").italic = True
        return doc

    def test_it_produces_the_blocks_the_service_produces(self) -> None:
        html = docx_html.to_html(self._doc())
        assert "<h1>A section</h1>" in html
        assert "<strong>Label: </strong>value" in html
        assert "<em>in italics</em>" in html

    def test_a_bullet_list_stays_one_block(self) -> None:
        """The service returns a list as a single <ul> chunk, not one chunk per item."""
        html = docx_html.to_html(self._doc())
        assert html.count("<ul>") == 1
        assert html.count("<li>") == 2

    def test_size_and_alignment_ride_along(self) -> None:
        html = docx_html.to_html(self._doc())
        assert 'style="text-align:center"' in html
        assert "font-size:18pt" in html

    def test_empty_paragraphs_are_dropped(self) -> None:
        """Word carries them for spacing; the service does not emit empty chunks."""
        html = docx_html.to_html(self._doc())
        assert "<p></p>" not in html
        assert "<p><span></span></p>" not in html
