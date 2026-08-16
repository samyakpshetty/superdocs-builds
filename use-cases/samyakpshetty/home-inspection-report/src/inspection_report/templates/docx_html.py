"""Turning a `.docx` format into the block HTML SuperDocs produces from the same file.

The binder works on HTML, because that is what comes back when a registered format is loaded
from SuperDocs. Two paths need the same HTML without the service:

* the deterministic fake, which backs the entire application with no API key;
* the fallback in :mod:`inspection_report.templates.registry`, so an inspector who has walked
  a property still gets their report when SuperDocs is unreachable.

So this converter has one job, and it is a fidelity job: **produce the same block structure
the live service produces**, not a prettier one. What the service returns for a `.docx` was
read off a real round trip — a flat sequence of ``<h1>`` / ``<p>`` / ``<ul>`` blocks, one per
chunk, with ``<strong>``, ``<em>`` and ``<span>`` inside them and lists kept whole. Anything
richer here would let a format work offline and fail live, which is the failure this build
has already been bitten by three times.
"""

from __future__ import annotations

import contextlib
import html as html_lib
from pathlib import Path
from typing import Any

from docx import Document


def _runs_to_html(paragraph: Any) -> str:
    out: list[str] = []
    for run in paragraph.runs:
        text = html_lib.escape(run.text)
        if not text:
            continue
        if run.bold:
            text = f"<strong>{text}</strong>"
        if run.italic:
            text = f"<em>{text}</em>"
        # Size and colour ride on a span, which is the shape the live service returns them
        # in. A format's letterhead is 18pt for a reason, and dropping it here would make
        # the offline document quietly plainer than the one the service produces.
        style = []
        if run.font.size is not None:
            style.append(f"font-size:{run.font.size.pt:g}pt")
        colour = getattr(run.font.color, "rgb", None) if run.font.color is not None else None
        if colour is not None:
            style.append(f"color:#{colour}")
        if style:
            text = f'<span style="{";".join(style)}">{text}</span>'
        out.append(text)
    return "".join(out)


def _paragraph_style(paragraph: Any) -> str:
    """Alignment and rules, as the inline style the service emits.

    The rules matter as much as the alignment: the formats close the masthead with a heavy
    border and sit every heading on a hairline, and a converter that dropped them produced
    HTML describing a plainer document than the `.docx` it came from.
    """
    from docx.oxml.ns import qn

    style: list[str] = []

    name = getattr(paragraph.alignment, "name", None)
    if name == "CENTER":
        style.append("text-align:center")
    elif name == "RIGHT":
        style.append("text-align:right")

    properties = paragraph._p.pPr
    borders = properties.find(qn("w:pBdr")) if properties is not None else None
    if borders is not None:
        for edge in ("bottom", "top"):
            border = borders.find(qn("w:" + edge))
            if border is None or border.get(qn("w:val")) in (None, "none", "nil"):
                continue
            # OOXML sizes a border in eighths of a point.
            eighths = border.get(qn("w:sz")) or "6"
            colour = (border.get(qn("w:color")) or "111318").lstrip("#")
            with contextlib.suppress(TypeError, ValueError):
                style.append(f"border-{edge}:{int(eighths) / 8:g}pt solid #{colour}")

    return f' style="{";".join(style)}"' if style else ""


def _heading_level(paragraph: Any) -> int | None:
    name = (paragraph.style.name or "") if paragraph.style is not None else ""
    if name.startswith("Heading "):
        try:
            return min(6, max(1, int(name.split()[-1])))
        except ValueError:
            return None
    if name == "Title":
        return 1
    return None


def _is_bullet(paragraph: Any) -> bool:
    name = (paragraph.style.name or "") if paragraph.style is not None else ""
    return name.startswith("List Bullet") or name.startswith("List Number")


def _table_to_html(table: Any) -> str:
    """A table, as one block.

    Tables carry the severity legend, and a converter that only walked `doc.paragraphs`
    silently dropped them — the offline document would have been missing a section the live
    one has, which is the fake-is-more-forgiving trap pointed the other way.
    """
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            inner = "".join(
                f"<p>{_runs_to_html(p)}</p>" for p in cell.paragraphs if _runs_to_html(p).strip()
            )
            style = _cell_style(cell)
            cells.append(f"<td{style}>{inner}</td>")
        rows.append(f"<tr>{''.join(cells)}</tr>")
    return f"<table>{''.join(rows)}</table>"


def _cell_style(cell: Any) -> str:
    """A cell's fill and width, as inline style.

    Both live in the OOXML and neither survives on its own: an exporter reading the HTML has
    no other source for them. Dropping the width is what turned a 4mm severity swatch into a
    third of the page — the colour carried and the geometry did not.
    """
    from docx.oxml.ns import qn

    properties = cell._tc.tcPr
    if properties is None:
        return ""

    style: list[str] = []
    shade = properties.find(qn("w:shd"))
    if shade is not None:
        fill = shade.get(qn("w:fill"))
        if fill and fill.lower() not in ("auto", "ffffff"):
            style.append(f"background-color:#{fill}")

    width = properties.find(qn("w:tcW"))
    if width is not None and width.get(qn("w:type")) == "dxa":
        with contextlib.suppress(TypeError, ValueError):
            # dxa is twentieths of a point; 1440 to the inch.
            style.append(f"width:{int(width.get(qn('w:w'))) / 1440:.2f}in")

    return f' style="{";".join(style)}"' if style else ""


def to_html(doc: Any) -> str:
    """Convert an open Word document to the flat block HTML the service returns.

    Walks the body in document order rather than `doc.paragraphs`, so a table appears where
    it actually sits rather than being dropped.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks: list[str] = []
    bullets: list[str] = []

    def flush() -> None:
        if bullets:
            items = "".join(f"<li>{b}</li>" for b in bullets)
            blocks.append(f"<ul>{items}</ul>")
            bullets.clear()

    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}tbl"):
            flush()
            blocks.append(_table_to_html(Table(child, doc)))
            continue
        if not child.tag.endswith("}p"):
            continue
        paragraph = Paragraph(child, doc)
        inner = _runs_to_html(paragraph)
        if not inner.strip():
            # Word documents carry empty paragraphs for spacing; the service drops them
            # rather than emitting empty chunks.
            continue
        if _is_bullet(paragraph):
            bullets.append(inner)
            continue
        flush()
        level = _heading_level(paragraph)
        align = _paragraph_style(paragraph)
        blocks.append(f"<h{level}{align}>{inner}</h{level}>" if level else f"<p{align}>{inner}</p>")
    flush()
    return "\n".join(blocks)


def from_path(path: Path) -> str:
    return to_html(Document(str(path)))


def from_bytes(data: bytes) -> str:
    import io

    return to_html(Document(io.BytesIO(data)))
