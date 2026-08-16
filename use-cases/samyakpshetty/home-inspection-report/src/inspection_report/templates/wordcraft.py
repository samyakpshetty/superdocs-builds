"""The Word typography python-docx does not expose.

A report format is the firm's document and it goes to a buyer, an agent and a lender. Built
from `add_paragraph` and the default styles it reads as something typed rather than something
designed — no rules, no letterspacing, no shading, no page furniture. All of that exists in
OOXML; python-docx simply has no API for it, so these are the small XML helpers that put it
back within reach.

Kept apart from :mod:`inspection_report.templates.authoring` so the formats themselves read
as a description of a document rather than as a pile of namespace calls.
"""

from __future__ import annotations

from typing import Any

from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Inches, Pt, RGBColor


def _element(tag: str, **attrs: str) -> Any:
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(f"w:{key}"), value)
    return node


def letterspace(run: Any, twentieths_of_a_point: int) -> None:
    """Track a run out. Small caps at normal tracking look cramped; this is what fixes it."""
    run._element.get_or_add_rPr().append(_element("w:spacing", val=str(twentieths_of_a_point)))


def small_caps(run: Any) -> None:
    run._element.get_or_add_rPr().append(_element("w:smallCaps", val="1"))


def rule(
    paragraph: Any,
    *,
    edge: str = "bottom",
    size: int = 6,
    colour: str = "111318",
    space: int = 8,
) -> None:
    """A border on one edge of a paragraph — the workhorse of a designed document.

    ``size`` is in eighths of a point, which is the unit OOXML uses: 6 is a hairline, 24 is
    a heavy rule.
    """
    borders = _element("w:pBdr")
    borders.append(
        _element("w:" + edge, val="single", sz=str(size), space=str(space), color=colour)
    )
    paragraph._p.get_or_add_pPr().append(borders)


def shade(cell: Any, colour: str) -> None:
    """Fill a table cell. Used for the severity swatches, and nowhere decorative."""
    cell._tc.get_or_add_tcPr().append(_element("w:shd", val="clear", fill=colour))


def cell_width(cell: Any, inches: float) -> None:
    """Set a cell's width.

    Through the property setter, not by appending a `w:tcW`: python-docx already writes one
    when the table is created, and a second is simply ignored — which is how a 4mm severity
    swatch silently stayed at the default third-of-the-page.
    """
    cell.width = Inches(inches)


def no_table_borders(table: Any) -> None:
    """A table used for layout should not look like a table."""
    borders = _element("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_element("w:" + edge, val="none", sz="0"))
    table._tbl.tblPr.append(borders)


def keep_with_next(paragraph: Any) -> None:
    """Stop a heading stranding itself at the foot of a page."""
    paragraph._p.get_or_add_pPr().append(_element("w:keepNext", val="1"))


def space_before(paragraph: Any, points: float) -> None:
    paragraph.paragraph_format.space_before = Pt(points)


def field(paragraph: Any, instruction: str) -> None:
    """Insert a Word field — this is how a footer says "page 2 of 9" and stays true."""
    begin = paragraph.add_run()
    begin._r.append(_element("w:fldChar", fldCharType="begin"))

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    holder = paragraph.add_run()
    holder._r.append(instr)

    end = paragraph.add_run()
    end._r.append(_element("w:fldChar", fldCharType="end"))


def page_footer(section: Any, *, left: str, ink: RGBColor) -> None:
    """A footer carrying what the document is, and where you are in it."""
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    rule(paragraph, edge="top", size=4, colour="C9C6BC", space=6)

    label = paragraph.add_run(left)
    label.font.size = Pt(7.5)
    label.font.color.rgb = ink
    small_caps(label)
    letterspace(label, 20)

    paragraph.add_run("\t")
    for text in ("Page ", "PAGE", " of ", "NUMPAGES"):
        if text.isupper() and text.isalpha():
            field(paragraph, text)
        else:
            run = paragraph.add_run(text)
            run.font.size = Pt(7.5)
            run.font.color.rgb = ink
